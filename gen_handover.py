"""
CHA Cambodia — Website Handover Guide generator.
Run: python gen_handover.py
Output: CHA_Website_Handover_Guide.docx
"""
import sys
import os

sys.path.insert(0, r"C:\Users\Asus\Desktop\Cha\.agents\skills\docgen\scripts")
from docgen_engine import (
    create_styled_document,
    add_executive_cover_page,
    add_heading_1,
    add_heading_2,
    add_heading_3,
    add_body_p,
    add_step_card,
    add_callout,
    add_styled_table,
    RGB_PRIMARY,
    RGB_SECONDARY,
    RGB_TEXT,
    RGB_MUTED,
    COLOR_PRIMARY,
    COLOR_SECONDARY,
    COLOR_SUCCESS,
    COLOR_WARNING,
    COLOR_SURFACE,
    COLOR_BORDER,
)
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\Asus\Desktop\Cha\CHA_Website_Handover_Guide.docx"
LOGO = r"C:\Users\Asus\Desktop\Cha\cha-cambodia-theme\logo-icon-cha.png"


def add_h1_bar(doc, text):
    """Heading 1 with thin bottom accent border."""
    p = add_heading_1(doc, text)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), COLOR_SECONDARY)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_num_steps(doc, steps):
    """Numbered steps as a simple styled list with navy numbers."""
    for i, s in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.line_spacing = 1.2
        r_n = p.add_run(f"{i}.  ")
        r_n.font.name = "Segoe UI"
        r_n.font.size = Pt(10)
        r_n.font.bold = True
        r_n.font.color.rgb = RGB_PRIMARY
        r_t = p.add_run(s)
        r_t.font.name = "Segoe UI"
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = RGB_TEXT


def add_bullets(doc, items):
    for s in items:
        add_body_p(doc, s, bullet=True)


def build():
    doc = create_styled_document()

    # Patch footer text for this document
    for section in doc.sections:
        fp = section.footer.paragraphs[0]
        fp.clear()
        frun = fp.add_run("CHA Website Handover Guide  |  Confidential  |  Page ")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGB_MUTED

    # ========== COVER ==========
    add_executive_cover_page(
        doc,
        title="Website Handover Guide",
        subtitle="Operations, content management, members, backups, and access transfer\nfor chacambodia.org — written for CHA leadership and staff",
        logo_path=LOGO if os.path.exists(LOGO) else None,
        meta_table={
            "Organization": "Cambodian Haemophilia Association (CHA Cambodia)",
            "Public Website": "https://chacambodia.org",
            "Admin Portal": "https://chacambodia.org/wp-admin",
            "Hosting": "Namecheap Stellar Business (LiteSpeed + cPanel)",
            "Document Version": "1.0.0 — September 2026",
            "Audience": "Executive Leadership, Operations & Content Managers",
            "Scope": "Website operations & accounts only (no mobile app)",
            "Status": "Official handover deliverable",
        },
    )

    # ========== 1. WELCOME ==========
    add_h1_bar(doc, "1. Welcome & How to Use This Guide")
    add_callout(
        doc,
        "This guide is the single handover pack for running the CHA Cambodia website day to day. "
        "It covers logging in, changing content without code, managing members, email, backups, "
        "and which accounts must be transferred to your team. Follow the step cards in order — "
        "after almost every change you must purge the LiteSpeed cache (Section 11).",
        title="Executive summary",
        alert_type="note",
    )
    add_heading_2(doc, "What this system is")
    add_bullets(doc, [
        "WordPress site on Namecheap Stellar Business hosting (LiteSpeed Enterprise + cPanel + CloudLinux).",
        "One shared MySQL database stores pages, news, campaigns, members, and donations.",
        "Custom theme: cha-cambodia-theme — content is edited through the WordPress admin, not code.",
        "Bilingual site (English + Khmer). Most text fields have an EN and a KM twin.",
        "Transactional email via Brevo SMTP from noreply@chacambodia.org.",
        "Offsite backups: UpdraftPlus → Google Drive (database daily, files weekly).",
    ])
    add_heading_2(doc, "How this guide is organised")
    add_styled_table(
        doc,
        ["Part", "Sections", "You will learn"],
        [
            ["Getting access", "2 – 3", "Which logins to receive; how to sign in to WordPress"],
            ["Editing the site", "4 – 6", "Customizer, News & Events, Campaigns"],
            ["Members & email", "7 – 8", "Add / edit / reset members; Brevo SMTP settings"],
            ["Safety net", "9 – 12", "Google Drive backups, restore, LiteSpeed, theme zip deploys"],
            ["Handover close", "13 – 16", "Developer boundary, credential checklist, open items, URLs"],
        ],
    )
    add_callout(
        doc,
        "Never share admin passwords in email or chat. Use a password manager or an in-person / "
        "video handover. Record each transferred account in the checklist in Section 14.",
        title="Golden rule for credentials",
        alert_type="important",
    )

    # ========== 2. ACCESS & LOGINS ==========
    add_h1_bar(doc, "2. Access & Logins — Handover Checklist")
    add_body_p(
        doc,
        "Before day-to-day work begins, the outgoing developer/agency must transfer every account below. "
        "Credentials are handed over out-of-band (password manager or spoken) — never inside this document or Git.",
    )
    add_heading_2(doc, "Accounts to transfer")
    add_styled_table(
        doc,
        ["Account", "Where it is used", "Priority", "Notes"],
        [
            ["WordPress admin", "chacambodia.org/wp-admin", "Essential", "All site editing; contains SMTP settings"],
            ["cPanel / Namecheap", "Hosting, files, DNS, SSL", "Essential", "Also needed for SPF DNS edit & full backups"],
            ["Google Drive (backups)", "UpdraftPlus remote storage", "Essential", "Hold the account that authorises Drive backups"],
            ["Brevo", "Email delivery / limits", "High", "300 emails/day free tier; monitor usage"],
            ["PayWay / ABA portal", "Donation gateway account", "If used", "Sandbox today; production credentials pending"],
            ["GitHub (repo)", "Source code for theme/app", "Optional", "Developer-facing; not required for daily content"],
        ],
    )
    add_heading_2(doc, "WordPress user roles")
    add_styled_table(
        doc,
        ["Role", "Give to", "Can do"],
        [
            ["Administrator", "Lead / IT owner", "Everything: themes, plugins, users, SMTP, settings"],
            ["Editor", "Comms staff", "Publish, edit, delete News, Campaigns, Pages"],
            ["Author", "Junior writers", "Write and publish their own articles + media"],
            ["Subscriber", "General public", "Read-only; used for the member portal only"],
        ],
    )
    add_callout(
        doc,
        "Only give Administrator to people who truly need full control. Everyone else should be Editor "
        "or Author. Review Users → All Users quarterly and remove leavers.",
        title="Least privilege",
        alert_type="important",
    )

    # ========== 3. LOGGING IN ==========
    add_h1_bar(doc, "3. Logging into WordPress — First Steps")
    add_step_card(doc, 1, "Open the admin portal", [
        "Browser: Chrome, Edge, Safari, or Firefox.",
        "Go to https://chacambodia.org/wp-admin and bookmark it.",
    ])
    add_step_card(doc, 2, "Sign in", [
        "Enter your username or email + password.",
        "Tick Remember Me only on a private, password-protected device.",
        "Click Log In to reach the Dashboard.",
    ])
    add_step_card(doc, 3, "First-session hygiene", [
        "If you cannot log in: click Lost your password? and follow the email link (check Spam).",
        "Change any password that was shared verbally during handover.",
        "Profile → update your display name and email if needed.",
    ])
    add_heading_2(doc, "Dashboard — what matters")
    add_bullets(doc, [
        "Posts / Pages / Media — standard WordPress content.",
        "News & Events and Campaigns — CHA-specific content (Sections 5–6).",
        "CHA Members — member roster, Add New, Edit, password reset, CSV export, SMTP (Section 7–8).",
        "Appearance → Customize — all homepage text, contacts, logos (Section 4).",
        "LiteSpeed Cache — toolbar icon for Purge All (Section 11).",
        "Settings, Tools, Plugins — Administrators only; avoid unless instructed.",
    ])
    add_heading_2(doc, "Media library hygiene")
    add_bullets(doc, [
        "Upload images < 2 MB where possible; use JPG/WebP for photos, PNG for logos.",
        "Recommended news featured images: about 1200 × 800 px.",
        "Give files real names (world-hemophilia-day-2026.jpg) before upload.",
        "Delete only unused media — never delete files still used by live pages.",
    ])

    # ========== 4. CUSTOMIZER ==========
    add_h1_bar(doc, "4. Customizing the Site Without Code (Customizer)")
    add_callout(
        doc,
        "Path: WordPress Admin → Appearance → Customize → edit → click Publish → then Purge LiteSpeed Cache "
        "(Section 11). Almost every text setting has an English field and a Khmer (ខ្មែរ) twin — fill both.",
        title="How to open the Customizer",
        alert_type="note",
    )
    add_heading_2(doc, "What you can change here")
    add_styled_table(
        doc,
        ["Panel / section", "Examples of editable content"],
        [
            ["Homepage Content → Hero", "Hero background image, title lines, subtitle, both button labels (EN+KM)"],
            ["Homepage Content → Statistics", "3 stat values + labels (Provinces, Patients, Partners)"],
            ["Homepage Content → How We Help", "Heading, sub, 4 help cards"],
            ["Homepage Content → CTA Banner", "Red donate banner heading, sub, button"],
            ["Homepage Content → About Section", "Who is CHA? heading, vision & mission, team image"],
            ["Contact & Footer → Contact Info", "Address, phone (display + dial digits), email, office hours"],
            ["Contact & Footer → Footer", "Tagline, copyright line"],
            ["Navigation & Header → Site Logo", "Header and footer logo images"],
            ["Navigation & Header → Menu Labels", "Up to 23 navigation labels (EN+KM)"],
            ["About / Programs / Haemophilia Page", "Full page copy, timeline, leaders, hospitals, medical cards"],
            ["Footer & Modals → Footer Navigation", "Quick links, resources, social, legal labels"],
            ["Footer & Modals → Donate / Member Modal", "Modal copy, amounts, login/register strings"],
            ["Homepage Extras", "Impact items, membership benefits, contact form labels"],
        ],
    )
    add_heading_2(doc, "Rules that prevent broken content")
    add_bullets(doc, [
        "Always click Publish — Draft changes are invisible to visitors.",
        "Phone: update BOTH fields (display number and digits used for tel: links).",
        "Do not clear a field unless you intend to hide that element.",
        "After Publish → Purge All (LiteSpeed) → hard-refresh the site (Ctrl+Shift+R).",
        "Brand colors (#0B1D6D navy, #E31E24 red, #6A2C91 purple) are NOT in the Customizer — developer change.",
    ])
    add_callout(
        doc,
        "If a section looks empty after your edit, purge cache first. If still wrong, use the Customizer "
        "Live Preview history / close without publishing and try again — Publish is the only save.",
        title="Tip",
        alert_type="note",
    )

    # ========== 5. NEWS ==========
    add_h1_bar(doc, "5. Publishing News & Events")
    add_body_p(
        doc,
        "News posts appear automatically as the three latest cards on the homepage and in the full "
        "archive at https://chacambodia.org/news (with category filters).",
    )
    add_heading_2(doc, "Create a new article")
    add_num_steps(doc, [
        "Left menu → News & Events → Add New Post.",
        "Enter a clear title (e.g. “World Hemophilia Day 2026 Workshop in Siem Reap”).",
        "Write the body in the editor; insert images with the + / Image block.",
        "In the right sidebar → Featured Image → Set featured image (≈1200×800).",
        "In the Article Details box set: Display Date (e.g. Apr 17, 2025), Category Badge "
        "(Event / Update / Workshop / Announcement), and optional Khmer title + excerpt.",
        "Optionally add a News Category in the right sidebar for archive filtering.",
        "Click Publish, then Purge LiteSpeed Cache and hard-refresh /news.",
    ])
    add_heading_2(doc, "Edit or unpublish")
    add_bullets(doc, [
        "News & Events → All Posts → hover → Edit.",
        "Update content or set Status to Draft / move to Trash to remove from the site.",
        "Always purge cache after edits so the homepage cards refresh.",
    ])
    add_callout(
        doc,
        "If the News list is empty, the homepage shows built-in demo/fallback cards. That does not mean "
        "your content failed — publish at least one real post and the live articles take over.",
        title="Fallback cards",
        alert_type="important",
    )

    # ========== 6. CAMPAIGNS ==========
    add_h1_bar(doc, "6. Managing Fundraising Campaigns")
    add_body_p(
        doc,
        "Campaigns appear in the homepage Donate tab as progress cards (up to four) and in the full list "
        "at https://chacambodia.org/campaigns. Progress % is calculated automatically from Raised ÷ Goal.",
    )
    add_heading_2(doc, "Add a campaign")
    add_num_steps(doc, [
        "Left menu → Campaigns → Add New Campaign.",
        "Enter the campaign title and description (what the funds support).",
        "In Campaign Details set: Icon, Raised Amount ($), Goal Amount ($), "
        "Theme Color (red / blue / purple), and optional Khmer title + description.",
        "Publish (or Update) → Purge LiteSpeed Cache.",
    ])
    add_heading_2(doc, "Update progress only")
    add_num_steps(doc, [
        "Campaigns → All Campaigns → click the campaign.",
        "Change Raised Amount (and Goal if the target changed).",
        "Click Update → purge cache. The homepage bar and % recalculate on next load.",
    ])
    add_callout(
        doc,
        "With zero published campaigns the homepage shows three built-in demo campaigns so the layout "
        "never looks broken. Publish real campaigns to replace them.",
        title="Demo fallback",
        alert_type="important",
    )
    add_callout(
        doc,
        "Corporate Partners (WFH / IFAH) under the campaign list is static content — developer edit only.",
        title="Note",
        alert_type="note",
    )

    # ========== 7. MEMBERS ==========
    add_h1_bar(doc, "7. Managing Members (Practical Tasks)")
    add_body_p(
        doc,
        "Left menu → CHA Members is the staff-only member database (separate from WP users). "
        "Use it to register people manually, fix profiles, reset passwords, and verify sign-ups.",
    )

    add_heading_2(doc, "7.1 Add a new member")
    add_num_steps(doc, [
        "CHA Members → click Add New (top of the page).",
        "Choose Membership Category: General Member or Patient "
        "(Patient reveals medical fields).",
        "General Information: Full Name (English) * and Email * are required; "
        "optionally Khmer name, phone, address.",
        "Password: leave blank to auto-generate, or type one for them (min. 6 characters recommended).",
        "If Patient: fill Date of Birth, Blood Type, Hemophilia Type / condition, and other medical fields.",
        "Optionally upload a photo (Add Photo).",
        "Save. The member appears in the list immediately.",
    ])
    add_callout(
        doc,
        "If you set a password yourself, give it to the member out-of-band and tell them to change it after "
        "first login. If you left Password blank, the system generated one — you can set a known password later via Edit.",
        title="Password handling",
        alert_type="note",
    )

    add_heading_2(doc, "7.2 Edit an existing member")
    add_num_steps(doc, [
        "CHA Members → All Members → search or filter → click Edit on the row.",
        "Update any field: name, Khmer name, email, phone, address, role/category, photo, medical details.",
        "Scroll to Set New Password only if you intend to change it — leave blank to keep the current password.",
        "Click Update / Save.",
    ])

    add_heading_2(doc, "7.3 Reset a member’s password")
    add_num_steps(doc, [
        "CHA Members → find the member → Edit.",
        "In General Information → Set New Password, type the new password "
        "(leave blank if you are not changing it).",
        "Save. The member signs in with the new password on the website member modal "
        "or in the registered flow associated with their email.",
    ])
    add_callout(
        doc,
        "Admin-side password set does not email the member automatically — contact them with the new "
        "credential securely, or ask them to use Forgot password? for a self-service email reset.",
        title="Important",
        alert_type="important",
    )

    add_heading_2(doc, "7.4 Verify a pending member")
    add_body_p(
        doc,
        "Members who register online stay Pending until they click the email verification link — or until "
        "staff verifies them manually. Unverified accounts cannot log in (they receive a “not verified” response).",
    )
    add_num_steps(doc, [
        "CHA Members → open the Pending Verification tab (or filter).",
        "Confirm the person’s identity (email / phone).",
        "Click Verify on their row.",
        "They can now log in normally.",
    ])

    add_heading_2(doc, "7.5 Delete a member & export CSV")
    add_bullets(doc, [
        "Delete: on the member row click Delete and confirm — this is permanent; prefer fixing the record instead.",
        "Export CSV: CHA Members → Export CSV for offline lists / reporting (names, emails, roles, etc.).",
        "Use the search box and role tabs (Patient, Family member, Healthcare professional, Member) to narrow lists.",
    ])

    # ========== 8. SMTP ==========
    add_h1_bar(doc, "8. Email (Brevo SMTP)")
    add_body_p(
        doc,
        "Verification emails, password-reset links, and admin notifications are sent through Brevo SMTP "
        "from noreply@chacambodia.org. Settings live inside the CHA Members screen — not in the public code.",
    )
    add_heading_2(doc, "Open the SMTP panel")
    add_num_steps(doc, [
        "WordPress Admin → CHA Members → click the SMTP button (top of the page).",
        "Review connection fields: host (smtp-relay.brevo.com), port (587 / TLS), username, "
        "Brevo SMTP key/password, from name and from email.",
        "Toggle the service enabled and optionally admin notification email.",
        "Save changes.",
        "Send a test (test email / SMTP test) and confirm receipt — check Spam if needed.",
    ])
    add_callout(
        doc,
        "Brevo free tier = 300 emails per day. Heavy verification/reset traffic can hit the cap — "
        "monitor the Brevo dashboard. SMTP debug logging writes to wp-content/cha-smtp-debug.log "
        "(developer should keep this non-public; currently an open hardening item).",
        title="Limits & security",
        alert_type="warning",
    )
    add_callout(
        doc,
        "Emails may land in Spam until DNS is fixed. The SPF record must include Brevo: "
        "add include:sendinblue.com to the domain SPF in Namecheap DNS. This is an open item (Section 15).",
        title="Deliverability (SPF)",
        alert_type="important",
    )

    # ========== 9. GOOGLE DRIVE ==========
    add_h1_bar(doc, "9. Connecting Google Drive for Backups")
    add_body_p(
        doc,
        "Destination: Google Drive — free tier is enough at CHA’s scale, one-click OAuth setup, "
        "restorable from the same plugin. Follow once during handover, then verify backups keep running.",
    )
    add_heading_2(doc, "9.1 Install UpdraftPlus")
    add_num_steps(doc, [
        "WP Admin → Plugins → Add New.",
        "Search UpdraftPlus (author: UpdraftPlus.Com) → Install Now → Activate.",
        "Skip if already installed.",
    ])
    add_heading_2(doc, "9.2 Connect to Google Drive")
    add_num_steps(doc, [
        "WP Admin → Settings → UpdraftPlus Backups.",
        "Open the Settings tab.",
        "Remote storage: tick Google Drive (free — not under Get Premium).",
        "Click Connect / Authorise Google Drive → a Google sign-in window opens.",
        "Sign in with the Google account that should hold backups, allow access.",
        "Return to WordPress and click Test — expect success.",
    ])
    add_heading_2(doc, "9.3 Schedule & retention (same page)")
    add_styled_table(
        doc,
        ["Setting", "Recommended / live value"],
        [
            ["Files schedule", "Weekly"],
            ["Database schedule", "Daily"],
            ["Keep this many file backups", "4 live (guide often suggests 12 — raise if you want ~3 months)"],
            ["Keep this many DB backups", "14 live (guide often suggests 30 — raise for ~1 month)"],
            ["Delete local copies after upload", "Yes"],
            ["Email reports", "Off, or a monitored inbox"],
        ],
    )
    add_body_p(doc, "Click Save Changes after editing any schedule or retention value.")
    add_heading_2(doc, "9.4 Verify the first backup (always)")
    add_num_steps(doc, [
        "Settings → UpdraftPlus → Backup Now.",
        "Tick database + all files → run.",
        "Open Existing Backups: both DB and files must show a green tick (uploaded).",
        "Open Google Drive → UpdraftPlus folder → confirm archive files exist.",
    ])
    add_callout(
        doc,
        "A backup is only real when it is green in WordPress AND visible in Drive. "
        "Re-run Backup Now if either is missing.",
        title="What “good” looks like",
        alert_type="success",
    )

    # ========== 10. BACKUPS & RECOVERY ==========
    add_h1_bar(doc, "10. Backups & Disaster Recovery (Owner Runbook)")
    add_heading_2(doc, "Normal operations")
    add_bullets(doc, [
        "Automated: UpdraftPlus → Google Drive (files weekly, database daily) — see Section 9.",
        "Before any theme zip or risky change: run a manual Backup Now (files + DB) and wait for green ticks.",
        "Hosting safety net: enable cPanel AutoBackup if your plan offers it (confirm once with Namecheap — open item).",
        "Code safety net: the development repo is pushed to GitHub with checkpoint tags after major phases.",
    ])
    add_heading_2(doc, "If something breaks — restore paths")
    add_styled_table(
        doc,
        ["Problem", "Restore path"],
        [
            ["Bad content / wrong copy", "Edit or trash the post/page in WP Admin; purge cache"],
            ["Database corruption / lost posts", "UpdraftPlus → Existing Backups → Restore → database only"],
            ["Theme / whole site broken", "UpdraftPlus full restore (files + DB), or re-upload last known good theme zip"],
            ["Server-level disaster", "cPanel full backup / host restore; phpMyAdmin import if needed"],
            ["Member photos missing", "Re-upload or restore /wp-content/uploads/cha-photos/"],
        ],
    )
    add_callout(
        doc,
        "Never restore from an untested backup. Once after handover (and then yearly), run a restore test "
        "on a safe copy or staging approach and confirm the site loads. Full step-by-step recovery detail "
        "lives in the project file BACKUP_AND_RESTORE.md (developer + owner copy).",
        title="Golden rule",
        alert_type="warning",
    )

    # ========== 11. LITESPEED ==========
    add_h1_bar(doc, "11. LiteSpeed Cache — The Golden Rule")
    add_callout(
        doc,
        "After ANY of these: theme zip upload · Customizer Publish · new/edited News or Campaign · "
        "plugin change · restored backup — click LiteSpeed Cache → Purge All in the top admin toolbar, "
        "then hard-refresh the public site (Ctrl+Shift+R / Cmd+Shift+R).",
        title="When to purge",
        alert_type="important",
    )
    add_heading_2(doc, "Exact steps")
    add_num_steps(doc, [
        "While logged into WP Admin, find the LiteSpeed Cache diamond icon in the top toolbar.",
        "Click it → Purge All.",
        "Open the public site in a clean/private window or hard-refresh.",
        "Confirm the change is visible; if not, purge once more and refresh.",
    ])
    add_body_p(
        doc,
        "Why it matters: LiteSpeed serves static cached pages for speed. Skipping a purge means visitors "
        "keep seeing old phone numbers, old campaigns, or a stale homepage — a failure mode already seen in production.",
    )

    # ========== 12. THEME UPDATE ==========
    add_h1_bar(doc, "12. Deploying a Theme Update (When a Developer Sends a Zip)")
    add_callout(
        doc,
        "Only Administrators do this. Never edit PHP/CSS/JS yourself unless you know the impact — "
        "ask the developer instead (Section 13).",
        title="Who should deploy",
        alert_type="important",
    )
    add_heading_2(doc, "Deploy steps")
    add_num_steps(doc, [
        "Backup first: UpdraftPlus → Backup Now (files + database) → green ticks.",
        "WordPress Admin → Appearance → Themes → Add New Theme → Upload Theme.",
        "Choose the cha-cambodia-theme.zip file → Install Now → Activate "
        "(or replace the active theme if your workflow uses update-in-place).",
        "Purge LiteSpeed Cache (Section 11) and hard-refresh.",
        "Verify: homepage, /news, /campaigns, member login modal, contact info, language toggle EN/KM.",
        "If anything is wrong, restore the previous backup / previous zip and purge again.",
    ])

    # ========== 13. DEVELOPER ==========
    add_h1_bar(doc, "13. When to Call a Developer")
    add_heading_2(doc, "Safe for content/staff (no code)")
    add_bullets(doc, [
        "Customizer text, images, contacts, logos, menu labels.",
        "News & Campaign posts, media library, pages content where the editor allows it.",
        "CHA Members add/edit/verify/password reset/CSV; SMTP settings panel.",
        "UpdraftPlus backups and restores; LiteSpeed purge; theme zip upload for Administrators.",
    ])
    add_heading_2(doc, "Developer-only (request a change)")
    add_styled_table(
        doc,
        ["Change", "Why it needs a developer"],
        [
            ["Brand colors / fonts / layout CSS", "Hard-coded in style-cha.css"],
            ["Bank account number or QR image filename", "Hard-coded in theme templates"],
            ["Social media URLs, hard-coded links", "In front-page / footer PHP"],
            ["Legal page body copy structure", "Inside dedicated page templates"],
            ["SMTP debug log hardening / rate limits", "functions.php security work"],
            ["Any functions.php / REST API change", "Backend code + zip re-deploy"],
            ["Newsletter plugins, new integrations", "Needs testing against existing theme"],
        ],
    )

    # ========== 14. CREDENTIALS CHECKLIST ==========
    add_h1_bar(doc, "14. Credentials Transfer Checklist (Sign-off)")
    add_body_p(
        doc,
        "Complete during the handover meeting. Leave blank until the credential has actually been moved "
        "into CHA’s password manager and tested.",
    )
    add_styled_table(
        doc,
        ["Asset", "Where it lives", "Transferred (✓)", "Date", "Notes / who holds it"],
        [
            ["WordPress admin (owner account)", "chacambodia.org/wp-admin", "", "", ""],
            ["cPanel / Namecheap login", "namecheap.com / hosting panel", "", "", ""],
            ["Google Drive (backup account)", "Google account OAuth for UpdraftPlus", "", "", ""],
            ["Brevo account", "app.brevo.com", "", "", ""],
            ["PayWay / ABA portal", "payway / sandbox portal", "", "", ""],
            ["DNS / SPF edit access", "Namecheap domain DNS", "", "", ""],
            ["GitHub repo (optional)", "ndxdigitalsupport/Cha_Cambodia", "", "", ""],
            ["Theme zip + this guide", "Shared drive / handover pack", "", "", ""],
            ["Test staff WP logins", "Users → All Users", "", "", ""],
        ],
    )

    # ========== 15. OPEN ITEMS ==========
    add_h1_bar(doc, "15. Open Items the Client Inherits")
    add_body_p(doc, "Website-only punch list still open at handover (September 2026):")
    add_styled_table(
        doc,
        ["#", "Item", "Owner", "Why it matters"],
        [
            ["1", "Add include:sendinblue.com to SPF (Namecheap DNS)", "CHA / DNS holder", "Email deliverability"],
            ["2", "Keep wp-content/cha-smtp-debug.log non-public", "Developer", "Security hardening"],
            ["3", "Confirm one manual UpdraftPlus Backup Now → green tick", "CHA admin", "Proves Drive pipeline"],
            ["4", "Confirm cPanel AutoBackup is enabled", "CHA / Namecheap", "Hosting safety net (Phase B)"],
            ["5", "Run a restore test from Google Drive backup", "CHA + developer", "Never trust untested backups"],
            ["6", "Confirm SR chapter / placeholder addresses & content", "CHA leadership", "Accuracy of Locations/content"],
            ["7", "Monitor Brevo 300 emails/day cap", "CHA admin", "Avoid dropped verification emails"],
            ["8", "Raise DB/file retention if 14/4 too short", "CHA admin", "Longer recovery window"],
            ["9", "WordPress users review (remove leavers)", "Administrator", "Access hygiene"],
        ],
    )

    # ========== 16. APPENDICES ==========
    add_h1_bar(doc, "16. Appendices")
    add_heading_2(doc, "A. Key URLs")
    add_styled_table(
        doc,
        ["Resource", "URL"],
        [
            ["Public website", "https://chacambodia.org"],
            ["WordPress admin", "https://chacambodia.org/wp-admin"],
            ["News archive", "https://chacambodia.org/news"],
            ["Campaigns archive", "https://chacambodia.org/campaigns"],
            ["REST API base", "https://chacambodia.org/wp-json/cha/v1"],
            ["Password reset (members)", "https://chacambodia.org/reset-password"],
            ["SMTP debug log (restrict!)", "https://chacambodia.org/wp-content/cha-smtp-debug.log"],
            ["Google Drive backups", "drive.google.com → UpdraftPlus folder"],
            ["Brevo", "https://app.brevo.com"],
            ["Namecheap / cPanel", "namecheap.com → panel"],
        ],
    )
    add_heading_2(doc, "B. Glossary")
    add_styled_table(
        doc,
        ["Term", "Meaning"],
        [
            ["Customizer", "WordPress screen (Appearance → Customize) for visual, no-code edits"],
            ["Slug", "The URL ending, e.g. /news or /campaigns"],
            ["CPT", "Custom Post Type — News & Events and Campaigns are CHA’s"],
            ["Purge cache", "Force LiteSpeed to forget old page copies (Purge All)"],
            ["Theme zip", "The deployable file a developer sends for site updates"],
            ["UpdraftPlus", "The backup plugin connected to Google Drive"],
            ["Brevo", "Email delivery service (formerly Sendinblue) used via SMTP"],
            ["KHQR", "Cambodian QR payment standard used for donations"],
            ["EN / KM twin fields", "Paired English and Khmer text inputs for bilingual copy"],
        ],
    )
    add_heading_2(doc, "C. Related documents in the project")
    add_styled_table(
        doc,
        ["Document", "Purpose"],
        [
            ["CHA_WEBSITE_ADMIN_GUIDE.md (+ Khmer)", "Longer operations manual (EN + KM)"],
            ["BACKUP_SETUP_GUIDE.md", "Detailed UpdraftPlus / Drive setup"],
            ["BACKUP_AND_RECOVERY_PLAN.md", "Backup phase plan + checkpoint log"],
            ["BACKUP_AND_RESTORE.md", "Emergency restore runbook"],
            ["DEPLOYMENT_CHECKLIST.md", "Pre-deploy backup gate for the developer"],
        ],
    )
    add_callout(
        doc,
        "End of guide. Questions during handover: walk through Sections 2–3 (access), 7 (members), "
        "and 9–11 (backups + cache) live with your team before the developer disconnects.",
        title="Handover close",
        alert_type="success",
    )

    doc.save(OUT)
    print(f"OK saved {OUT}")


if __name__ == "__main__":
    build()
